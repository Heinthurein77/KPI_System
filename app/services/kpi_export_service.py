"""Excel rendering for individual monthly KPI reports.

This service is deliberately read-only: it turns existing submission records
and the already-calculated KPI summary into a downloadable workbook without
changing scores, templates, workflows, or database data.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.models.kpi_submission import KPISubmission
from app.models.user import User


_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_SUMMARY_FILL = PatternFill("solid", fgColor="EAF2F8")
_ALTERNATE_FILL = PatternFill("solid", fgColor="F8FAFC")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="1E3A5F")
_LABEL_FONT = Font(bold=True, color="334155")
_THIN_BORDER = Border(bottom=Side(style="thin", color="CBD5E1"))


def _write_summary_cell(sheet: Worksheet, label_cell: str, value_cell: str, label: str, value: object) -> None:
    sheet[label_cell] = label
    sheet[label_cell].font = _LABEL_FONT
    sheet[value_cell] = value


def build_employee_kpi_report(
    employee: User,
    submissions: list[KPISubmission],
    year: int,
    month: str,
    combined_score: dict | None,
) -> BytesIO:
    """Return a formatted .xlsx report for one employee and monthly period."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPI Report"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:H1")
    sheet["A1"] = "Employee KPI Score Report"
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].alignment = Alignment(horizontal="left")

    _write_summary_cell(sheet, "A3", "B3", "Employee", employee.name)
    _write_summary_cell(sheet, "A4", "B4", "Email", employee.email)
    _write_summary_cell(sheet, "A5", "B5", "Department", employee.department.name if employee.department else "—")
    _write_summary_cell(sheet, "D3", "E3", "Month", month)
    _write_summary_cell(sheet, "D4", "E4", "Year", year)

    total_weight = sum(float(submission.kpi_template.weight) for submission in submissions)
    final_attainment = combined_score["attainment"] if combined_score else "N/A"
    final_status = combined_score["status"] if combined_score else "Not finalized"
    _write_summary_cell(sheet, "D5", "E5", "Total Weight", total_weight)
    _write_summary_cell(sheet, "G3", "H3", "Final KPI Score", final_attainment)
    _write_summary_cell(sheet, "G4", "H4", "KPI Status", final_status)
    sheet["E5"].number_format = "0.00"
    if isinstance(final_attainment, (int, float)):
        sheet["H3"].number_format = "0.00"

    headers = [
        "KPI Metric",
        "Weight",
        "Target",
        "Self Score",
        "Dept Score",
        "Final Score",
        "Actual Score",
        "Status",
    ]
    for column_number, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column_number, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER

    for row_number, submission in enumerate(submissions, start=8):
        actual_score = submission.effective_score
        values = [
            submission.kpi_template.metric_name,
            float(submission.kpi_template.weight),
            float(submission.kpi_template.target),
            submission.self_score,
            submission.dept_score,
            submission.final_score,
            actual_score,
            submission.status.value,
        ]
        for column_number, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column_number, value=value)
        if row_number % 2 == 0:
            for cell in sheet[row_number]:
                cell.fill = _ALTERNATE_FILL
        for cell in sheet[row_number]:
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center")
        for column in "BCDEFG":
            sheet[f"{column}{row_number}"].number_format = "0.00"

    summary_range = sheet["A3:H5"]
    for row in summary_range:
        for cell in row:
            cell.fill = _SUMMARY_FILL
            cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A8"
    widths = {"A": 34, "B": 12, "C": 12, "D": 14, "E": 14, "F": 14, "G": 14, "H": 24}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _write_annual_month_sheet(
    sheet: Worksheet,
    employee: User,
    submissions: list[KPISubmission],
    year: int,
    month: str,
    combined_score: dict | None,
) -> None:
    """Write one month of detail inside an employee's annual workbook."""
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"KPI Score Detail — {month} {year}"
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].alignment = Alignment(horizontal="left")

    _write_summary_cell(sheet, "A3", "B3", "Employee", employee.name)
    _write_summary_cell(sheet, "A4", "B4", "Department", employee.department.name if employee.department else "—")
    _write_summary_cell(sheet, "D3", "E3", "Period", f"{month} {year}")
    _write_summary_cell(sheet, "D4", "E4", "Metrics", len(submissions))
    _write_summary_cell(
        sheet,
        "G3",
        "H3",
        "KPI Attainment",
        combined_score["attainment"] if combined_score else "N/A",
    )
    _write_summary_cell(sheet, "G4", "H4", "Status", combined_score["status"].strip() if combined_score else "Not finalized")
    if combined_score:
        sheet["H3"].number_format = "0.00"

    headers = [
        "KPI Metric",
        "Weight",
        "Target",
        "Self Score",
        "Dept Score",
        "Final Score",
        "Actual Score",
        "Status",
    ]
    for column_number, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column_number, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER

    if not submissions:
        sheet.merge_cells("A8:H8")
        sheet["A8"] = "No KPI records for this month."
        sheet["A8"].alignment = Alignment(horizontal="center")
        sheet["A8"].fill = _ALTERNATE_FILL
    else:
        for row_number, submission in enumerate(submissions, start=8):
            values = [
                submission.kpi_template.metric_name,
                float(submission.kpi_template.weight),
                float(submission.kpi_template.target),
                submission.self_score,
                submission.dept_score,
                submission.final_score,
                submission.effective_score,
                submission.status.value,
            ]
            for column_number, value in enumerate(values, start=1):
                sheet.cell(row=row_number, column=column_number, value=value)
            if row_number % 2 == 0:
                for cell in sheet[row_number]:
                    cell.fill = _ALTERNATE_FILL
            for cell in sheet[row_number]:
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(vertical="center")
            for column in "BCDEFG":
                sheet[f"{column}{row_number}"].number_format = "0.00"

    for row in sheet["A3:H4"]:
        for cell in row:
            cell.fill = _SUMMARY_FILL
            cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A8"
    widths = {"A": 34, "B": 12, "C": 12, "D": 14, "E": 14, "F": 14, "G": 14, "H": 24}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def build_employee_annual_kpi_report(
    employee: User,
    submissions: list[KPISubmission],
    year: int,
    months: list[str],
    monthly_scores: dict[str, dict | None],
) -> BytesIO:
    """Return one workbook containing an employee's KPI data for a full year.

    The first sheet is a year-at-a-glance summary; each subsequent month keeps
    the same metric-level detail as the normal monthly report.
    """
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Annual Summary"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:H1")
    summary["A1"] = "Employee Annual KPI Score Report"
    summary["A1"].font = _TITLE_FONT
    summary["A1"].alignment = Alignment(horizontal="left")

    grouped = {month: [] for month in months}
    for submission in submissions:
        grouped[submission.month_or_quarter].append(submission)

    finalized = [score["attainment"] for score in monthly_scores.values() if score]
    annual_average = round(sum(finalized) / len(finalized), 2) if finalized else "N/A"
    _write_summary_cell(summary, "A3", "B3", "Employee", employee.name)
    _write_summary_cell(summary, "A4", "B4", "Email", employee.email)
    _write_summary_cell(summary, "A5", "B5", "Department", employee.department.name if employee.department else "—")
    _write_summary_cell(summary, "D3", "E3", "Year", year)
    _write_summary_cell(summary, "D4", "E4", "Months with KPI data", sum(bool(grouped[m]) for m in months))
    _write_summary_cell(summary, "G3", "H3", "Annual average", annual_average)
    _write_summary_cell(summary, "G4", "H4", "Finalized months", len(finalized))
    if isinstance(annual_average, (int, float)):
        summary["H3"].number_format = "0.00"

    headers = ["Month", "Metrics", "Finalized", "KPI Attainment", "Status"]
    for column_number, header in enumerate(headers, start=1):
        cell = summary.cell(row=7, column=column_number, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER

    for row_number, month in enumerate(months, start=8):
        score = monthly_scores[month]
        values = [
            month,
            len(grouped[month]),
            score["scored_count"] if score else 0,
            score["attainment"] if score else None,
            score["status"].strip() if score else "No finalized KPI scores",
        ]
        for column_number, value in enumerate(values, start=1):
            summary.cell(row=row_number, column=column_number, value=value)
        if row_number % 2 == 0:
            for cell in summary[row_number]:
                cell.fill = _ALTERNATE_FILL
        for cell in summary[row_number]:
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center")
        summary[f"D{row_number}"].number_format = "0.00"

    for row in summary["A3:H5"]:
        for cell in row:
            cell.fill = _SUMMARY_FILL
            cell.alignment = Alignment(vertical="center")
    summary.freeze_panes = "A8"
    for column, width in {"A": 18, "B": 14, "C": 14, "D": 18, "E": 28}.items():
        summary.column_dimensions[column].width = width

    for month in months:
        detail = workbook.create_sheet(month)
        _write_annual_month_sheet(detail, employee, grouped[month], year, month, monthly_scores[month])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
